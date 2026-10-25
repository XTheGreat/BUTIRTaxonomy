import io
import os
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

OPSI = ("A", "B", "C", "D")

_SIMBOL = {
    r"\cdot": "·", r"\times": "×", r"\div": "÷", r"\pm": "±", r"\mp": "∓",
    r"\leq": "≤", r"\le": "≤", r"\geq": "≥", r"\ge": "≥", r"\neq": "≠", r"\ne": "≠",
    r"\approx": "≈", r"\equiv": "≡", r"\infty": "∞", r"\rightarrow": "→", r"\to": "→",
    r"\leftarrow": "←", r"\Rightarrow": "⇒", r"\leftrightarrow": "↔", r"\rightleftharpoons": "⇌",
    r"\degree": "°", r"\circ": "°", r"\%": "%", r"\ldots": "...", r"\dots": "...",
    r"\alpha": "α", r"\beta": "β", r"\gamma": "γ", r"\delta": "δ", r"\epsilon": "ε",
    r"\theta": "θ", r"\lambda": "λ", r"\mu": "μ", r"\pi": "π", r"\rho": "ρ", r"\sigma": "σ",
    r"\tau": "τ", r"\phi": "φ", r"\omega": "ω", r"\Delta": "Δ", r"\Omega": "Ω", r"\Sigma": "Σ",
    r"\sum": "Σ", r"\int": "∫", r"\partial": "∂", r"\nabla": "∇",
}
_YUNANI = {
    r"\alpha", r"\beta", r"\gamma", r"\delta", r"\epsilon", r"\theta", r"\lambda", r"\mu", r"\pi",
    r"\rho", r"\sigma", r"\tau", r"\phi", r"\omega", r"\Delta", r"\Omega", r"\Sigma",
}
_SIMBOL_URUT = sorted(_SIMBOL, key=len, reverse=True)


def _hapus_pembungkus(teks: str, perintah: str) -> str:
    return re.sub(rf"\\{perintah}\{{([^{{}}]*)\}}", r"\1", teks)


def latex_ke_teks(teks: str | None) -> str:
    if not teks:
        return ""
    t = teks.replace("\\\\(", "\\(").replace("\\\\)", "\\)")
    t = re.sub(r"\\[()\[\]]", "", t)
    t = t.replace("$$", "").replace("$", "")

    for perintah in ("ce", "text", "mathrm", "mathbf", "mathit", "textbf", "textit", "boldsymbol"):
        t = _hapus_pembungkus(t, perintah)

    for _ in range(3):
        t = re.sub(r"\\d?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", t)
        t = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", t)
    t = re.sub(r"\\sqrt\s*(\w)", r"√\1", t)

    for kunci in _SIMBOL_URUT:
        spasi = r"[ ]?(?=[A-Za-z]\b)" if kunci in _YUNANI else ""
        t = re.sub(re.escape(kunci) + r"(?![A-Za-z])" + spasi, _SIMBOL[kunci], t)

    t = t.replace("<=>", "⇌").replace("->", "→").replace("<-", "←")

    t = re.sub(r"\^\{([^{}]*)\}", lambda m: _pangkat(m.group(1)), t)
    t = re.sub(r"_\{([^{}]*)\}", lambda m: _indeks(m.group(1)), t)
    t = re.sub(r"\^(\d+[+-]?|[+-]|\w)", lambda m: _pangkat(m.group(1)), t)
    t = re.sub(r"_(\w)", lambda m: _indeks(m.group(1)), t)

    t = re.sub(r"\\(left|right|big|Big|displaystyle)\b", "", t)
    t = re.sub(r"\\[,;:! ]", " ", t)
    t = t.replace("\\\\", " ")
    t = re.sub(r"\\([A-Za-z]+)", r"\1", t)
    t = t.replace("{", "").replace("}", "")
    return re.sub(r"[ \t]+", " ", t).strip()


_SUPER = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
_SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")


def _pangkat(isi: str) -> str:
    return isi.translate(_SUPER) if all(c in "0123456789+-=()n" for c in isi) else f"^({isi})"


def _indeks(isi: str) -> str:
    return isi.translate(_SUB) if all(c in "0123456789+-=()" for c in isi) else f"_({isi})"


MEDIA_DOKUMEN = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


async def buat_dokumen(
    format: str, judul: str, butir_list: list, dengan_kunci: bool = True, subjudul: str | None = None
) -> tuple[bytes, str]:
    import asyncio

    pembuat = {"pdf": buat_pdf, "docx": buat_docx}[format]
    isi = await asyncio.to_thread(pembuat, judul, butir_list, dengan_kunci, subjudul)
    return isi, MEDIA_DOKUMEN[format]


_KANDIDAT_FONT = [
    ("C:/Windows/Fonts/times.ttf", "C:/Windows/Fonts/timesbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/dejavu/DejaVuSerif-Bold.ttf"),
    ("/Library/Fonts/Times New Roman.ttf", "/Library/Fonts/Times New Roman Bold.ttf"),
]
_font_cache: tuple[str, str, bool] | None = None


def _font_pdf() -> tuple[str, str, bool]:
    global _font_cache
    if _font_cache:
        return _font_cache
    for reguler, tebal in _KANDIDAT_FONT:
        if os.path.exists(reguler) and os.path.exists(tebal):
            try:
                pdfmetrics.registerFont(TTFont("NaskahReg", reguler))
                pdfmetrics.registerFont(TTFont("NaskahTebal", tebal))
                _font_cache = ("NaskahReg", "NaskahTebal", True)
                return _font_cache
            except Exception:
                continue
    _font_cache = ("Times-Roman", "Times-Bold", False)
    return _font_cache


def _bersih_pdf(teks: str, unicode_ok: bool) -> str:
    teks = escape(teks).replace("\n", "<br/>")
    if not unicode_ok:
        teks = teks.encode("latin-1", "replace").decode("latin-1")
    return teks


def buat_pdf(judul: str, butir_list: list, dengan_kunci: bool = True, subjudul: str | None = None) -> bytes:
    reguler, tebal, uni = _font_pdf()
    dasar = ParagraphStyle("dasar", fontName=reguler, fontSize=11, leading=15)
    gaya_judul = ParagraphStyle("judul", parent=dasar, fontName=tebal, fontSize=16, leading=20, alignment=TA_CENTER)
    gaya_sub = ParagraphStyle("sub", parent=dasar, fontSize=10, alignment=TA_CENTER, textColor="#5A6070")
    gaya_soal = ParagraphStyle("soal", parent=dasar, spaceBefore=10, spaceAfter=4, leftIndent=18, firstLineIndent=-18)
    gaya_opsi = ParagraphStyle("opsi", parent=dasar, leftIndent=36, firstLineIndent=-16, spaceAfter=1)
    gaya_kunci = ParagraphStyle("kunci", parent=dasar, fontSize=10, leftIndent=18, spaceBefore=4, textColor="#1F6B3A")
    gaya_bahas = ParagraphStyle("bahas", parent=dasar, fontSize=10, leftIndent=18, spaceBefore=2, textColor="#5A6070")

    def p(teks: str, gaya) -> Paragraph:
        return Paragraph(_bersih_pdf(teks, uni), gaya)

    isi = [p(judul, gaya_judul)]
    if subjudul:
        isi.append(p(subjudul, gaya_sub))
    isi += [Spacer(1, 6), HRFlowable(width="100%", thickness=0.6, color="#8C8778"), Spacer(1, 4)]

    for i, b in enumerate(butir_list, 1):
        blok = [p(f"{i}.  {latex_ke_teks(b.pertanyaan)}", gaya_soal)]
        pilihan = b.pilihan or {}
        for k in OPSI:
            blok.append(p(f"{k}.  {latex_ke_teks(pilihan.get(k, ''))}", gaya_opsi))
        if dengan_kunci:
            blok.append(
                p("   ·   ".join(x for x in (f"Kunci: {b.kunci_jawaban}", b.level_bloom, f"Bobot {b.bobot_poin}") if x), gaya_kunci)
            )
            if b.pembahasan:
                blok.append(p(f"Pembahasan: {latex_ke_teks(b.pembahasan)}", gaya_bahas))
        isi.append(KeepTogether(blok))

    if not butir_list:
        isi.append(p("Tidak ada butir soal.", dasar))

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=judul, author="BUTIR",
    ).build(isi)
    return buf.getvalue()


def buat_docx(judul: str, butir_list: list, dengan_kunci: bool = True, subjudul: str | None = None) -> bytes:
    doc = Document()
    for bagian in doc.sections:
        bagian.left_margin = bagian.right_margin = Cm(2.2)
        bagian.top_margin = bagian.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(judul)
    run.bold = True
    run.font.size = Pt(16)
    if subjudul:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.add_run(subjudul).font.size = Pt(10)

    for i, b in enumerate(butir_list, 1):
        soal = doc.add_paragraph()
        soal.paragraph_format.space_before = Pt(10)
        soal.paragraph_format.space_after = Pt(3)
        soal.paragraph_format.keep_with_next = True
        soal.add_run(f"{i}.  ").bold = True
        soal.add_run(latex_ke_teks(b.pertanyaan))

        pilihan = b.pilihan or {}
        for k in OPSI:
            o = doc.add_paragraph()
            o.paragraph_format.left_indent = Cm(1.1)
            o.paragraph_format.space_after = Pt(1)
            o.paragraph_format.keep_with_next = dengan_kunci or k != "D"
            o.add_run(f"{k}.  {latex_ke_teks(pilihan.get(k, ''))}")

        if dengan_kunci:
            kn = doc.add_paragraph()
            kn.paragraph_format.left_indent = Cm(0.6)
            kn.paragraph_format.space_before = Pt(3)
            r = kn.add_run("   ·   ".join(x for x in (f"Kunci: {b.kunci_jawaban}", b.level_bloom, f"Bobot {b.bobot_poin}") if x))
            r.font.size = Pt(10)
            r.bold = True
            if b.pembahasan:
                pb = doc.add_paragraph()
                pb.paragraph_format.left_indent = Cm(0.6)
                rb = pb.add_run(f"Pembahasan: {latex_ke_teks(b.pembahasan)}")
                rb.font.size = Pt(10)

    if not butir_list:
        doc.add_paragraph("Tidak ada butir soal.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
